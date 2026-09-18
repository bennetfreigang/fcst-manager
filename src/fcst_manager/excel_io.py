"""Excel-Ein-/Ausgabe fuer Dateien mit beliebig vielen Artikelzeilen.

Erwartetes Layout (wie in den Referenzdateien):
  Zeile 1 = Abschnittsname je Spalte ("Historie" / "Order" / "FCST", leer bei Metaspalten)
  Zeile 2 = Spaltenkopf (Metafeldname bzw. Monat "JJJJ_MM")
  Zeile 3 ff = je Zeile ein Artikel
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import IO, TypeAlias

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .model import Decision, Item, clean_number
from .periods import Month, month_range

Source: TypeAlias = "str | Path | IO[bytes]"
"""Pfad oder offener Byte-Stream - letzteres fuer Uploads (Streamlit)."""

SECTION_ROW = 1
HEADER_ROW = 2
FIRST_DATA_ROW = 3

_META_FIELDS = {
    "item_number": ("itemnumber", "item number", "artikelnummer", "item"),
    "avg_demand": ("avg demand", "demand"),
    "moq": ("moq",),
    "lt": ("lt", "lead time", "leadtime"),
}


def _norm(value: object) -> str:
    return str(value).strip().lower() if value is not None else ""


def _resolve_meta_columns(headers: dict[int, object]) -> tuple[dict[str, int], list[str]]:
    """Ordnet Metaspalten robust zu (Gross-/Kleinschreibung, Zusaetze wie '[mon]')."""
    warnings: list[str] = []
    found: dict[str, int] = {}
    for field_name, needles in _META_FIELDS.items():
        for col, header in headers.items():
            h = _norm(header)
            if not h:
                continue
            if any(h == n or h.startswith(n) or n in h for n in needles):
                if field_name not in found:
                    found[field_name] = col
                break
    for required in ("item_number", "lt"):
        if required not in found:
            warnings.append(f"Pflichtspalte fuer '{required}' nicht gefunden")
    for optional in ("avg_demand", "moq"):
        if optional not in found:
            warnings.append(f"Spalte fuer '{optional}' fehlt -> wird als unbekannt behandelt")
    return found, warnings


@dataclass
class Layout:
    """Spaltenaufteilung eines Arbeitsblatts."""

    sheet_name: str
    meta: dict[str, int]
    historie: dict[Month, int]
    order: dict[Month, int]
    fcst: dict[Month, int]
    data_rows: list[int]
    warnings: list[str] = field(default_factory=list)

    @property
    def fcst_months(self) -> list[Month]:
        return sorted(self.fcst)

    def stichtag(self) -> Month:
        """Erste FCST-Spalte = 'aktueller Zeitpunkt' (so vom Kunden festgelegt)."""
        if not self.fcst:
            raise ValueError(
                "Kein FCST-Abschnitt gefunden - Stichtag nicht ableitbar. "
                "Bitte --stichtag JJJJ_MM angeben."
            )
        return self.fcst_months[0]


def read_layout(path: Source, sheet: str | None = None) -> tuple[Layout, openpyxl.Workbook]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]

    sections = {c: ws.cell(SECTION_ROW, c).value for c in range(1, ws.max_column + 1)}
    headers = {c: ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)}

    meta, warnings = _resolve_meta_columns(
        {c: h for c, h in headers.items() if not _norm(sections[c])}
    )

    buckets: dict[str, dict[Month, int]] = {"historie": {}, "order": {}, "fcst": {}}
    for col, section in sections.items():
        key = _norm(section)
        if key not in buckets:
            continue
        try:
            month = Month.parse(headers[col])
        except ValueError as exc:
            warnings.append(f"Spalte {col} im Abschnitt '{section}' uebersprungen: {exc}")
            continue
        if month in buckets[key]:
            warnings.append(
                f"Monat {month} kommt im Abschnitt '{section}' mehrfach vor - "
                f"Spalte {col} gewinnt"
            )
        buckets[key][month] = col

    item_col = meta.get("item_number")
    data_rows = [
        r
        for r in range(FIRST_DATA_ROW, ws.max_row + 1)
        if item_col and ws.cell(r, item_col).value not in (None, "")
    ]

    layout = Layout(
        sheet_name=ws.title,
        meta=meta,
        historie=buckets["historie"],
        order=buckets["order"],
        fcst=buckets["fcst"],
        data_rows=data_rows,
        warnings=warnings,
    )
    return layout, wb


def _read_series(ws, row: int, cols: dict[Month, int], label: str) -> tuple[dict[Month, float], list[str]]:
    series: dict[Month, float] = {}
    warnings: list[str] = []
    for month, col in cols.items():
        value, warn = clean_number(ws.cell(row, col).value, field_name=f"{label} {month}")
        if warn:
            warnings.append(warn)
        series[month] = value if value is not None else 0.0
    return series, warnings


def read_items(path: Source, sheet: str | None = None):
    """Liest alle Artikelzeilen. Liefert (Layout, Items, manuelle FCST-Werte je Item)."""
    layout, wb = read_layout(path, sheet)
    ws = wb[layout.sheet_name]

    items: list[Item] = []
    manual: dict[str, dict[Month, float]] = {}

    for row in layout.data_rows:
        warnings: list[str] = []
        item_number = str(ws.cell(row, layout.meta["item_number"]).value).strip()

        avg_demand = moq = None
        if "avg_demand" in layout.meta:
            avg_demand, warn = clean_number(
                ws.cell(row, layout.meta["avg_demand"]).value, field_name="AVG Demand"
            )
            if warn:
                warnings.append(warn)
        if "moq" in layout.meta:
            moq, warn = clean_number(ws.cell(row, layout.meta["moq"]).value, field_name="MOQ")
            if warn:
                warnings.append(warn)

        lt_raw, warn = clean_number(
            ws.cell(row, layout.meta["lt"]).value if "lt" in layout.meta else None,
            field_name="LT",
        )
        if warn:
            warnings.append(warn)
        if lt_raw is None:
            warnings.append("LT fehlt -> 0 angenommen; Anchor stuetzt sich nur auf die Order")
            lt = 0
        elif lt_raw < 0:
            warnings.append(f"LT {lt_raw:g} ist negativ -> 0 angenommen")
            lt = 0
        else:
            lt = int(round(lt_raw))
            if lt != lt_raw:
                warnings.append(f"LT {lt_raw:g} auf {lt} Monate gerundet")

        if avg_demand is not None and avg_demand <= 0:
            warnings.append(f"AVG Demand {avg_demand:g} <= 0 -> als unbekannt behandelt")
            avg_demand = None
        if moq is not None and moq <= 0:
            warnings.append(f"MOQ {moq:g} <= 0 -> als unbekannt behandelt")
            moq = None

        historie, w1 = _read_series(ws, row, layout.historie, "Historie")
        order, w2 = _read_series(ws, row, layout.order, "Order")
        manual_fcst, w3 = _read_series(ws, row, layout.fcst, "FCST")
        warnings.extend(w1 + w2 + w3)

        if not historie:
            warnings.append("kein Historie-Abschnitt gefunden")

        overlap = set(historie) & set(order)
        if overlap:
            warnings.append(
                f"Monate in Historie UND Order: {sorted(m.label for m in overlap)}"
            )

        items.append(
            Item(
                item_number=item_number,
                avg_demand=avg_demand,
                moq=moq,
                lt=lt,
                historie=historie,
                order=order,
                row=row,
                warnings=warnings,
            )
        )
        manual[item_number] = {m: q for m, q in manual_fcst.items() if q}

    wb.close()
    return layout, items, manual


# ---------------------------------------------------------------------------
# Getrennte Uploads (SalesHistorie / OrderBook / Artikelstammdaten)
# ---------------------------------------------------------------------------
#
# Alternative zur Sammeldatei oben: statt einer Datei mit Historie/Order/FCST
# in einem Blatt liefert der Nutzer drei separate, einfachere Dateien mit nur
# einer Kopfzeile. Diese werden zu einer Sammeldatei im Standardlayout
# zusammengefuehrt (build_combined_workbook) - ab da laufen read_items,
# write_output und die gesamte App unveraendert weiter.


def _read_item_series_table(
    path: Source, sheet: str | None = None
) -> tuple[dict[str, dict[Month, float]], list[str]]:
    """Liest eine Datei mit einer Kopfzeile: ItemNumber + Monatsspalten (JJJJ_MM).

    Fuer SalesHistorie- und OrderBook-Upload - im Gegensatz zur Sammeldatei gibt
    es hier nur einen Abschnitt je Datei, eine Abschnittszeile ist daher nicht
    noetig.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    headers = {c: ws.cell(1, c).value for c in range(1, ws.max_column + 1)}

    item_col: int | None = None
    month_cols: dict[Month, int] = {}
    warnings: list[str] = []
    for col, header in headers.items():
        h = _norm(header)
        if not h:
            continue
        if item_col is None and any(h == n or h.startswith(n) or n in h for n in _META_FIELDS["item_number"]):
            item_col = col
            continue
        try:
            month = Month.parse(header)
        except ValueError:
            warnings.append(f"Spalte {col} ('{header}') ist weder ItemNumber- noch Monatsspalte - ignoriert")
            continue
        if month in month_cols:
            warnings.append(f"Monat {month} kommt mehrfach vor - Spalte {col} gewinnt")
        month_cols[month] = col

    if item_col is None:
        warnings.append("Pflichtspalte fuer 'item_number' nicht gefunden")
        wb.close()
        return {}, warnings

    result: dict[str, dict[Month, float]] = {}
    for row in range(2, ws.max_row + 1):
        raw_number = ws.cell(row, item_col).value
        if raw_number in (None, ""):
            continue
        item_number = str(raw_number).strip()
        series: dict[Month, float] = {}
        for month, col in month_cols.items():
            value, warn = clean_number(ws.cell(row, col).value, field_name=f"{item_number} {month}")
            if warn:
                warnings.append(warn)
            if value:
                series[month] = value
        result[item_number] = series

    wb.close()
    return result, warnings


def read_sales_history_table(path: Source, sheet: str | None = None):
    """SalesHistorie-Upload: ItemNumber + Monatsspalten mit Verkaufsmengen."""
    return _read_item_series_table(path, sheet)


def read_orderbook_table(path: Source, sheet: str | None = None):
    """OrderBook-Upload: ItemNumber + Monatsspalten mit offenen Bestellmengen."""
    return _read_item_series_table(path, sheet)


def read_old_fcst_table(path: Source, sheet: str | None = None):
    """Upload eines frueheren FCST-Laufs zum Vergleich: ItemNumber + Monatsspalten
    (gleiches einfaches Format wie SalesHistorie/OrderBook). Unabhaengig vom
    Eingabemodus nutzbar - auch im Split-Upload, dessen zusammengefuehrte Datei
    selbst keine befuellten FCST-Spalten hat (siehe build_combined_workbook)."""
    return _read_item_series_table(path, sheet)


def read_item_master_table(
    path: Source, sheet: str | None = None
) -> tuple[dict[str, dict[str, float | int | None]], list[str]]:
    """Artikelstammdaten-Upload: ItemNumber, AVG Demand, MOQ, LT (eine Kopfzeile)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    headers = {c: ws.cell(1, c).value for c in range(1, ws.max_column + 1)}

    meta_cols, warnings = _resolve_meta_columns(headers)
    if "item_number" not in meta_cols:
        wb.close()
        return {}, warnings

    result: dict[str, dict[str, float | int | None]] = {}
    for row in range(2, ws.max_row + 1):
        raw_number = ws.cell(row, meta_cols["item_number"]).value
        if raw_number in (None, ""):
            continue
        item_number = str(raw_number).strip()
        entry: dict[str, float | int | None] = {}
        for field_name in ("avg_demand", "moq", "lt"):
            col = meta_cols.get(field_name)
            if col is None:
                continue
            value, warn = clean_number(ws.cell(row, col).value, field_name=f"{item_number} {field_name}")
            if warn:
                warnings.append(warn)
            entry[field_name] = value
        result[item_number] = entry

    wb.close()
    return result, warnings


def build_combined_workbook(
    historie: dict[str, dict[Month, float]],
    order: dict[str, dict[Month, float]],
    meta: dict[str, dict[str, float | int | None]],
    item_numbers: list[str],
    stichtag: Month,
    horizon_months: int,
) -> openpyxl.Workbook:
    """Fuegt die drei getrennten Uploads zu einer Sammeldatei im Standardlayout
    zusammen (Zeile 1 Abschnitt, Zeile 2 Spaltenkopf, ab Zeile 3 je Artikel eine
    Zeile) - danach ist sie fuer read_items/write_output nicht mehr von einer
    manuell gepflegten Sammeldatei zu unterscheiden. Die FCST-Spalten bleiben
    leer; ihre Breite (Stichtag..Stichtag+horizon-1) legt fest, wie weit
    write_output spaeter beim Export befuellen kann.
    """
    hist_months = sorted({m for series in historie.values() for m in series})
    order_months = sorted({m for series in order.values() for m in series})
    fcst_months = month_range(stichtag, stichtag + (horizon_months - 1))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabelle1"

    meta_headers = ["ItemNumber", "AVG Demand", "MOQ", "LT"]
    sections: list[object] = [None] * len(meta_headers)
    sections += ["Historie"] * len(hist_months)
    sections += ["Order"] * len(order_months)
    sections += ["FCST"] * len(fcst_months)
    headers = (
        meta_headers
        + [m.label for m in hist_months]
        + [m.label for m in order_months]
        + [m.label for m in fcst_months]
    )

    ws.append(sections)
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True)

    for number in item_numbers:
        item_meta = meta.get(number, {})
        h = historie.get(number, {})
        o = order.get(number, {})
        row = [
            number,
            item_meta.get("avg_demand"),
            item_meta.get("moq"),
            item_meta.get("lt"),
        ]
        row += [h.get(m) for m in hist_months]
        row += [o.get(m) for m in order_months]
        row += [None] * len(fcst_months)
        ws.append(row)

    ws.freeze_panes = "E3"
    return wb


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

REPORT_COLUMNS = [
    ("Zeile", lambda d, i: i.row),
    ("ItemNumber", lambda d, i: d.item_number),
    ("Segment", lambda d, i: str(d.segment)),
    ("Segment-Begruendung", lambda d, i: d.segment_reason),
    ("Ast", lambda d, i: str(d.branch)),
    ("Ast-Begruendung", lambda d, i: d.branch_reason),
    ("Bestellungen", lambda d, i: d.n_orders),
    ("Historie-Fenster [mon]", lambda d, i: d.history_window_months),
    ("AVG Demand", lambda d, i: i.avg_demand),
    ("MOQ", lambda d, i: i.moq),
    ("LT [mon]", lambda d, i: i.lt),
    ("Letzte Bindung", lambda d, i: d.last_commitment.label if d.last_commitment else None),
    ("Art der Bindung", lambda d, i: d.last_commitment_kind),
    ("Lieferluecke", lambda d, i: d.gap_date.label if d.gap_date else None),
    ("Luecke ab Stichtag [mon]", lambda d, i: d.gap_months_from_stichtag),
    ("Grosse Luecke", lambda d, i: "ja" if d.large_gap else "nein"),
    ("Q", lambda d, i: d.qty),
    ("Q-Herleitung", lambda d, i: d.qty_source),
    ("T [mon]", lambda d, i: d.interval),
    ("T-Herleitung", lambda d, i: d.interval_source),
    ("impliziter Demand", lambda d, i: round(d.implied_demand, 2) if d.implied_demand else None),
    ("Anchor", lambda d, i: d.anchor.label if d.anchor else None),
    ("Anchor-Herleitung", lambda d, i: d.anchor_source),
    ("FCST-Termine", lambda d, i: len(d.fcst)),
    ("FCST-Monate", lambda d, i: ", ".join(p.label for p in d.fcst)),
    ("FCST-Summe", lambda d, i: sum(p.qty for p in d.fcst) or None),
    ("Abweichung FCST zu Demand", lambda d, i: d.demand_deviation),
    ("uebersprungen (Order vorhanden)", lambda d, i: ", ".join(m.label for m in d.skipped_months)),
    ("Annahmen", lambda d, i: " | ".join(d.assumptions)),
    ("Warnungen", lambda d, i: " | ".join(d.warnings)),
]


def write_output(
    source: Source,
    target: Source,
    layout: Layout,
    items: list[Item],
    decisions: dict[str, Decision],
    manual: dict[str, dict[Month, float]] | None = None,
) -> None:
    """Schreibt eine Kopie der Eingabe mit befuellten FCST-Spalten + Report-Blatt."""
    wb = openpyxl.load_workbook(source)  # ohne data_only -> Formatierung/Formeln bleiben
    ws = wb[layout.sheet_name]

    for item in items:
        decision = decisions.get(item.item_number)
        if decision is None:
            continue
        series = {p.month: p.qty for p in decision.fcst}
        for month, col in layout.fcst.items():
            ws.cell(item.row, col).value = series.get(month)

    _write_report_sheet(wb, items, decisions, manual or {})
    wb.save(target)
    wb.close()


def _write_report_sheet(wb, items, decisions, manual) -> None:
    for name in ("FCST-Report", "FCST-Abweichungen"):
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet("FCST-Report")
    compare = bool(any(manual.values()))
    headers = [name for name, _ in REPORT_COLUMNS]
    if compare:
        headers += ["Vergleich alter FCST", "Abweichende Monate"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "C2"

    for item in items:
        decision = decisions.get(item.item_number)
        if decision is None:
            continue
        row = [getter(decision, item) for _, getter in REPORT_COLUMNS]
        if compare:
            verdict, diffs = compare_with_manual(decision, manual.get(item.item_number, {}))
            row += [verdict, diffs]
        ws.append(row)

    widths = {1: 7, 2: 16, 3: 13, 4: 46, 5: 44, 6: 46}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def compare_with_manual(decision: Decision, manual: dict[Month, float]) -> tuple[str, str]:
    """Vergleicht den berechneten FCST mit einem vorhandenen manuellen FCST."""
    if not manual:
        return "kein manueller FCST", ""
    computed = {p.month: p.qty for p in decision.fcst}
    months = sorted(set(manual) | set(computed))
    diffs = [
        f"{m.label}: manuell {manual.get(m, 0):g} / berechnet {computed.get(m, 0):g}"
        for m in months
        if (manual.get(m, 0) or 0) != (computed.get(m, 0) or 0)
    ]
    if not diffs:
        return f"exakte Uebereinstimmung ({len(computed)} Termine)", ""
    hits = sum(1 for m in months if (manual.get(m, 0) or 0) == (computed.get(m, 0) or 0) and m in manual)
    return f"{hits}/{len(manual)} Monate getroffen", " | ".join(diffs)
