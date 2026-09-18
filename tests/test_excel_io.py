"""Excel-Ein-/Ausgabe und Regression gegen die Referenzdateien des Kunden."""

import openpyxl
import pytest

from fcst_manager.cli import main as cli_main
from fcst_manager.engine import forecast
from fcst_manager.excel_io import (
    build_combined_workbook,
    compare_with_manual,
    read_item_master_table,
    read_items,
    read_orderbook_table,
    read_sales_history_table,
    write_output,
)
from fcst_manager.model import Config
from fcst_manager.periods import Month

from .conftest import DATA

SAMPLE = DATA / "Sammeldatei_Test.xlsx"


def _wb(rows: list[list[object]], tmp_path, name="in.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


# --- Referenzdateien ------------------------------------------------------


def test_layout_derives_stichtag_from_first_fcst_column(reference_file):
    layout, items, _ = read_items(reference_file)
    assert layout.stichtag().label == "2026_08"
    assert len(layout.historie) == 42 and len(layout.order) == 14 and len(layout.fcst) == 18
    assert len(items) == 1


def test_meta_columns_are_resolved_including_dot_placeholders(reference_file):
    _, items, _ = read_items(reference_file)
    item = items[0]
    assert item.lt == 5
    assert (item.avg_demand is None) == (item.item_number == "D228025-100")
    assert (item.moq is None) == (item.item_number == "D228025-100")


def test_reference_forecast_matches_manual_exactly(reference_file):
    """D228025-100 (Example.xlsx) muss Monat fuer Monat reproduziert werden - der
    Anchor trifft hier zufaellig auch mit der aktuellen Stichtag+LT-Formel.

    1/136648 (Example 2.xlsx) trifft NICHT mehr exakt: die Datei nennt Anchor
    2027_03, die vom Kunden definitiv bestaetigte Formel (Stichtag+LT, siehe
    engine.compute_anchor) ergibt 2027_01. Laut Kunde koennen die per Hand
    erstellten Referenzwerte selbst fehlerhaft sein - Q und T (die von der
    Anchor-Aenderung unberuehrt sind) muessen trotzdem weiter stimmen.
    """
    layout, items, manual = read_items(reference_file)
    decision = forecast(items[0], layout.stichtag(), Config())
    verdict, diffs = compare_with_manual(decision, manual[items[0].item_number])

    if items[0].item_number == "1/136648":
        assert decision.qty == 150 and decision.interval == 3
        assert "2027_03: manuell 150 / berechnet 0" in diffs, (
            "bekannte, vom Kunden akzeptierte Abweichung zum Anchor - falls diese "
            "Assertion bricht, hat sich entweder die Formel geaendert (dann Kommentar "
            "und Test anpassen) oder die Referenzdatei wurde per Hand korrigiert"
        )
        return

    assert diffs == ""
    assert verdict.startswith("exakte Uebereinstimmung")


# --- Sammeldatei mit vielen Zeilen ---------------------------------------


def test_sample_workbook_reads_all_rows():
    layout, items, _ = read_items(SAMPLE)
    assert len(items) == 12 == len(layout.data_rows)
    assert {i.item_number for i in items} >= {"D228025-100", "1/136648", "SLEEP-001"}


def test_unreadable_values_become_warnings_not_crashes():
    _, items, _ = read_items(SAMPLE)
    bad = next(i for i in items if i.item_number == "EDGE-BAD-010")
    assert bad.avg_demand is None and bad.moq is None and bad.lt == 0
    assert any("unlesbarer Wert" in w for w in bad.warnings)
    assert any("LT fehlt" in w for w in bad.warnings)


def test_negative_history_is_preserved_for_reporting():
    _, items, _ = read_items(SAMPLE)
    neg = next(i for i in items if i.item_number == "EDGE-NEG-009")
    assert min(neg.historie.values()) == -30
    assert neg.sale_events(), "negative Monate zaehlen nicht als Bestellung"
    assert len(neg.sale_events()) == 3


def test_export_writes_fcst_columns_and_report_sheet(tmp_path):
    layout, items, manual = read_items(SAMPLE)
    stichtag = layout.stichtag()
    decisions = {i.item_number: forecast(i, stichtag, Config()) for i in items}
    target = tmp_path / "out.xlsx"
    write_output(SAMPLE, target, layout, items, decisions, manual)

    wb = openpyxl.load_workbook(target)
    assert "FCST-Report" in wb.sheetnames
    ws = wb[layout.sheet_name]
    for item in items:
        series = decisions[item.item_number].as_series()
        for month, col in layout.fcst.items():
            assert (ws.cell(item.row, col).value or 0) == series.get(month.label, 0)
    report = wb["FCST-Report"]
    assert report.max_row == len(items) + 1
    assert "Q-Herleitung" in [c.value for c in report[1]]


def test_export_leaves_source_file_untouched(tmp_path):
    before = SAMPLE.read_bytes()
    layout, items, manual = read_items(SAMPLE)
    decisions = {i.item_number: forecast(i, layout.stichtag(), Config()) for i in items}
    write_output(SAMPLE, tmp_path / "out.xlsx", layout, items, decisions, manual)
    assert SAMPLE.read_bytes() == before


def test_missing_fcst_section_demands_explicit_stichtag(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, "Historie"])
    ws.append(["ItemNumber", "2025_01"])
    ws.append(["X", 10])
    path = tmp_path / "ohne_fcst.xlsx"
    wb.save(path)
    layout, _, _ = read_items(path)
    with pytest.raises(ValueError, match="--stichtag"):
        layout.stichtag()


# --- CLI ------------------------------------------------------------------


def test_cli_run_produces_output_and_log(tmp_path, capsys):
    out, log = tmp_path / "out.xlsx", tmp_path / "audit.txt"
    rc = cli_main(["run", str(SAMPLE), "-o", str(out), "--log", str(log)])
    assert rc == 0 and out.exists() and log.exists()
    printed = capsys.readouterr().out
    assert "Verteilung auf die Aeste" in printed
    text = log.read_text(encoding="utf-8")
    assert "Item D228025-100" in text and "Intervall T" in text


def test_cli_stichtag_override_shifts_the_forecast(tmp_path, capsys):
    for flag in (["--stichtag", "2026_08"], ["--stichtag", "2027_02"]):
        cli_main(["run", str(SAMPLE), "-o", str(tmp_path / "o.xlsx"), *flag])
    assert "Stichtag  : 2027_02" in capsys.readouterr().out


def test_cli_explain_filters_by_item(capsys):
    assert cli_main(["explain", str(SAMPLE), "--item", "1/136648"]) == 0
    printed = capsys.readouterr().out
    assert "Item 1/136648" in printed and "Item SLEEP-001" not in printed


def test_cli_explain_reports_unknown_item(capsys):
    assert cli_main(["explain", str(SAMPLE), "--item", "GIBTSNICHT"]) == 1


def test_cli_compare_t_lists_only_items_without_demand(capsys):
    assert cli_main(["compare-t", str(SAMPLE)]) == 0
    printed = capsys.readouterr().out
    assert "D228025-100" in printed and "1/136648" not in printed


def test_cli_rejects_unknown_t_method():
    with pytest.raises(SystemExit):
        cli_main(["run", str(SAMPLE), "--t-method", "gibtsnicht"])


# --- Getrennte Uploads: SalesHistorie / OrderBook / Artikelstammdaten -----


def test_read_sales_history_table_reads_one_header_row(tmp_path):
    path = _wb(
        [
            ["ItemNumber", "2025_01", "2025_02"],
            ["A", 10, 0],
            ["B", ".", 5],
        ],
        tmp_path,
    )
    historie, warnings = read_sales_history_table(path)
    assert warnings == []
    assert historie["A"] == {Month.parse("2025_01"): 10.0}
    assert historie["B"] == {Month.parse("2025_02"): 5.0}


def test_read_orderbook_table_ignores_unparseable_extra_columns(tmp_path):
    path = _wb(
        [
            ["ItemNumber", "2026_07", "Kommentar"],
            ["A", 40, "Split-Lieferung"],
        ],
        tmp_path,
    )
    order, warnings = read_orderbook_table(path)
    assert order["A"] == {Month.parse("2026_07"): 40.0}
    assert any("weder ItemNumber- noch Monatsspalte" in w for w in warnings)


def test_read_item_master_table_resolves_tolerant_columns_and_dot_placeholder(tmp_path):
    path = _wb(
        [
            ["Item Number", "AVG Demand", "MOQ", "LT [mon]"],
            ["A", 12.5, 50, 5],
            ["B", ".", ".", 3],
        ],
        tmp_path,
    )
    meta, warnings = read_item_master_table(path)
    assert warnings == []
    assert meta["A"] == {"avg_demand": 12.5, "moq": 50.0, "lt": 5.0}
    assert meta["B"] == {"avg_demand": None, "moq": None, "lt": 3.0}


def test_build_combined_workbook_roundtrips_through_read_items(tmp_path):
    """Simuliert den App-Merge-Schritt: drei getrennte Quellen -> Sammeldatei ->
    read_items() sieht danach eine ganz normale Sammeldatei."""
    historie = {
        "A": {Month.parse("2025_01"): 10.0, Month.parse("2025_06"): 10.0},
        "B": {Month.parse("2025_03"): 5.0},
    }
    order = {"A": {Month.parse("2026_09"): 10.0}}
    meta = {"A": {"avg_demand": 12.5, "moq": 5.0, "lt": 3.0}}
    stichtag = Month.parse("2026_08")

    wb = build_combined_workbook(historie, order, meta, ["A", "B"], stichtag, horizon_months=6)
    path = tmp_path / "merged.xlsx"
    wb.save(path)

    layout, items, _ = read_items(path)
    assert layout.stichtag() == stichtag
    assert len(layout.fcst) == 6
    by_number = {i.item_number: i for i in items}

    a = by_number["A"]
    assert a.avg_demand == 12.5 and a.moq == 5.0 and a.lt == 3
    assert a.historie[Month.parse("2025_01")] == 10.0
    assert a.order[Month.parse("2026_09")] == 10.0

    b = by_number["B"]
    assert b.avg_demand is None and b.moq is None
    assert "LT fehlt" in " ".join(b.warnings), "B hat keinen Metadaten-Eintrag -> LT unbekannt"
    assert b.historie[Month.parse("2025_03")] == 5.0
    assert b.order.get(Month.parse("2026_09"), 0) == 0, "B war nicht im OrderBook -> keine Order"

    decision = forecast(a, stichtag, Config())
    assert decision.segment is not None  # laeuft durch die normale Engine ohne Sonderfall
