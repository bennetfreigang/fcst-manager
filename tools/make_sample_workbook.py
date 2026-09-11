"""Erzeugt eine Sammeldatei mit vielen Artikelzeilen zum Testen der CLI.

Enthaelt die beiden vom Kunden validierten Zeilen plus Kantenfaelle, die je einen
Ast des Entscheidungsbaums bzw. eine Datenstoerung abdecken.

    uv run python tools/make_sample_workbook.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from fcst_manager.periods import Month, month_range

HIST = month_range(Month.parse("2023_01"), Month.parse("2026_06"))   # 42 Monate
ORDER = month_range(Month.parse("2026_07"), Month.parse("2027_08"))  # 14 Monate
FCST = month_range(Month.parse("2026_08"), Month.parse("2028_01"))   # 18 Monate

# (ItemNumber, AvgDemand, MOQ, LT, Historie, Order, Kommentar)
ROWS = [
    ("D228025-100", ".", ".", 5,
     {"2024_11": 20, "2024_12": 20, "2025_09": 40, "2025_10": 40, "2025_11": 20, "2026_03": 40},
     {"2026_09": 40, "2026_10": 40},
     "Referenz 1: High Runner, kleine Luecke, AVG Demand unbekannt"),
    ("1/136648", 37.861111111111114, 50, 5,
     {"2023_03": 300, "2024_06": 200, "2024_12": 200, "2025_05": 150,
      "2025_07": 150, "2025_08": 150, "2025_12": 150, "2026_05": 150},
     {"2026_09": 150, "2026_12": 150},
     "Referenz 2: High Runner, kleine Luecke, AVG Demand bekannt"),
    ("SLEEP-001", ".", ".", 4, {"2024_01": 10}, {},
     "Sleeper: eine alte Bestellung, kein Demand"),
    ("NEU-002", ".", ".", 6, {}, {},
     "Sleeper: gar keine Historie und keine Order"),
    ("MID-003", 20, 30, 3, {"2025_10": 60, "2026_02": 60}, {},
     "Mid Runner, kleine Luecke -> Standard (Annahme, nicht validiert)"),
    ("MIDGAP-004", 15, 50, 6, {"2025_04": 90, "2026_01": 90}, {"2027_06": 90},
     "Mid Runner, grosse Luecke, Backlog + Demand vorhanden -> FCST"),
    ("MIDGAP-005", ".", 50, 6, {"2025_04": 90, "2026_01": 90}, {"2027_06": 90},
     "Mid Runner, grosse Luecke, Backlog aber kein Demand -> kein FCST"),
    ("HIGHGAP-006", 25, 100, 9,
     {"2024_02": 200, "2024_09": 200, "2025_04": 200, "2025_11": 200, "2026_04": 200},
     {"2027_05": 200},
     "High Runner, grosse Luecke -> Menge aus Demand"),
    ("HIGHGAP-007", ".", 100, 9,
     {"2024_02": 200, "2024_09": 200, "2025_04": 200, "2025_11": 200, "2026_04": 200},
     {"2027_05": 200},
     "High Runner, grosse Luecke, kein Demand -> kein FCST"),
    ("EDGE-LT0-008", ".", ".", 0, {"2026_01": 5}, {},
     "Kantenfall: LT=0 + eine Bestellung (frueher Endlosschleife)"),
    ("EDGE-NEG-009", ".", ".", 4,
     {"2025_03": -30, "2025_06": 30, "2025_11": 30, "2026_02": 30}, {},
     "Kantenfall: negative Historie-Menge"),
    ("EDGE-BAD-010", "keine Angabe", ".", None, {"2025_06": 70, "2026_01": 70}, {"2026_11": 70},
     "Kantenfall: unlesbarer Demand, MOQ '.', LT leer"),
]


def build(target: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabelle1"

    meta = ["ItemNumber", "AVG Demand ÖBB [mon]", "MOQ Vertrag", "LT [mon]"]
    sections = [None] * len(meta) + ["Historie"] * len(HIST) + ["Order"] * len(ORDER) + ["FCST"] * len(FCST)
    headers = meta + [m.label for m in HIST + ORDER + FCST]
    ws.append(sections)
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True)

    for number, demand, moq, lt, hist, order, _comment in ROWS:
        row = [number, demand, moq, lt]
        row += [hist.get(m.label, 0) for m in HIST]
        row += [order.get(m.label, 0) for m in ORDER]
        row += [None] * len(FCST)          # FCST wird von der Engine befuellt
        ws.append(row)

    notes = wb.create_sheet("Testfaelle")
    notes.append(["ItemNumber", "abgedeckter Fall"])
    notes["A1"].font = notes["B1"].font = Font(bold=True)
    for number, *_, comment in ROWS:
        notes.append([number, comment])
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 70

    ws.freeze_panes = "E3"
    wb.save(target)
    print(f"{target}  ({len(ROWS)} Artikel, {len(HIST)}/{len(ORDER)}/{len(FCST)} Monate)")


if __name__ == "__main__":
    build(Path(__file__).resolve().parents[1] / "src" / "data" / "Sammeldatei_Test.xlsx")
